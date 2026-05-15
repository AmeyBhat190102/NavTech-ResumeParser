import io
from docx import Document

from app.parsers.base import BaseParser


class DOCXParser(BaseParser):
    """Extracts text from .docx files using python-docx."""

    def extract_text(self, file_bytes: bytes) -> str:
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]

        # Also pull text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text:
                        paragraphs.append(cell_text)

        return "\n".join(paragraphs).strip()
